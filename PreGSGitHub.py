import os
import time
import random
import numpy as np
import pandas as pd
import torch
import torch.nn as nn
import torch.nn.functional as F
from torch_geometric.data import Data
from torch_geometric.nn import GATConv, SAGEConv
from torch_geometric.utils import add_remaining_self_loops
from datetime import datetime
from openpyxl.styles import PatternFill, Font
from sklearn.metrics import f1_score

# ---------- 通用配置（固定超参，无调优） ----------
device = torch.device(
    'cuda' if torch.cuda.is_available() else 'cpu'
)

seed = 42


def set_random_seed(seed_value):
    """
    统一设置Python、NumPy、CPU和GPU随机种子。
    """
    random.seed(seed_value)
    np.random.seed(seed_value)
    torch.manual_seed(seed_value)

    if torch.cuda.is_available():
        torch.cuda.manual_seed_all(seed_value)


set_random_seed(seed)

if torch.cuda.is_available():
    torch.backends.cudnn.deterministic = True
    torch.backends.cudnn.benchmark = False

# 固定超参数（原论文GAT超参）
FIXED_LR = 0.005
FIXED_WD = 5e-4
FIXED_DROPOUT = 0.6
FIXED_EPOCHS = 2000
FIXED_PATIENCE = 100

# GS聚合器配置（2个mean、3个max、3个add）
FIXED_GS_AGGRS = ['mean', 'mean', 'max', 'max', 'max', 'add', 'add', 'add']

# 维度配置
FIXED_MLP_HIDDEN_DIM = 256
FIXED_GAT_HIDDEN_DIM = 8
FIXED_GAT_HEADS = [8, 1]

# 数据集比例配置
TRAIN_PERCENT_LIST = [0.2]
VAL_PERCENT = 0.1

# 模型名称常量
MODEL_GAT = "GAT"
MODEL_PREGS = "PreGS"
MODEL_PREGS_V2 = "PreGSv2"

# 所有可用模型的固定展示顺序

MODEL_ORDER = [
    MODEL_GAT,
    MODEL_PREGS,
    MODEL_PREGS_V2,
]

SELECTED_MODELS = [
    MODEL_GAT,
    MODEL_PREGS,
    MODEL_PREGS_V2,
]

def get_active_models():
    """
    检查本次选择的模型，并按照 MODEL_ORDER 的固定顺序返回。

    注意：
    PreGS 和 PreGSv2 都需要预训练 GAT 参数。
    因此，只要选择 PreGS 或 PreGSv2，就必须同时选择 GAT。
    """
    if not SELECTED_MODELS:
        raise ValueError("SELECTED_MODELS 不能为空，至少选择一个模型。")

    unknown_models = [
        model_name
        for model_name in SELECTED_MODELS
        if model_name not in MODEL_ORDER
    ]

    if unknown_models:
        raise ValueError(
            f"SELECTED_MODELS 中存在未知模型：{unknown_models}"
        )

    if len(SELECTED_MODELS) != len(set(SELECTED_MODELS)):
        raise ValueError(
            "SELECTED_MODELS 中存在重复模型，请删除重复项。"
        )

    needs_gat = (
        MODEL_PREGS in SELECTED_MODELS
        or MODEL_PREGS_V2 in SELECTED_MODELS
    )

    if needs_gat and MODEL_GAT not in SELECTED_MODELS:
        raise ValueError(
            "PreGS 和 PreGSv2 依赖预训练 GAT。"
            "选择 PreGS 或 PreGSv2 时，必须同时选择 MODEL_GAT。"
        )

    # 无论用户在 SELECTED_MODELS 中怎样排列，
    # 最终都按照 MODEL_ORDER 的固定顺序运行。
    return [
        model_name
        for model_name in MODEL_ORDER
        if model_name in SELECTED_MODELS
    ]

output_dir = os.path.join(
    os.path.dirname(os.path.abspath(__file__)),
    "output_gat_vs_pregs"
)
os.makedirs(output_dir, exist_ok=True)

# 累计保存所有数据集、训练比例和模型的统计结果
all_results_raw = []


def compute_macro_f1(y_true, y_pred):
    """Compute Macro-F1 on CPU without affecting training."""
    return f1_score(
        y_true.detach().cpu().numpy(),
        y_pred.detach().cpu().numpy(),
        average='macro',
        zero_division=0
    )


def compute_sample_std(values):
    """
    计算多次重复实验的样本标准差。

    使用 ddof=1：
    - 更适合论文中报告多次独立实验的波动；
    - 当只有一次实验时直接返回 0，避免产生 NaN。
    """
    values = np.asarray(values, dtype=np.float64)

    if values.size <= 1:
        return 0.0

    return float(np.std(values, ddof=1))


def format_mean_std(mean_value, std_value, scale=1.0, decimals=2):
    """
    将均值和标准差格式化为：
    84.25 ± 1.36

    参数：
    scale=100：将 0～1 的准确率和 F1 转为百分比；
    scale=1：用于运行时间等原始数值。
    """
    mean_value = float(mean_value) * scale
    std_value = float(std_value) * scale

    return (
        f"{mean_value:.{decimals}f} "
        f"± {std_value:.{decimals}f}"
    )


# ---------- 辅助函数：为表格前三名标色 ----------
def highlight_top3(sheet, start_row, start_col, end_row, end_col):
    """
    为指定区域内每行的前三名数值单元格标色
    :param sheet: openpyxl的worksheet对象
    :param start_row: 起始行（包含表头，从1开始）
    :param start_col: 起始列（从1开始）
    :param end_row: 结束行
    :param end_col: 结束列
    """
    # 定义颜色（第一名：深绿，第二名：浅绿，第三名：淡绿）
    fill_1st = PatternFill(start_color="00CC00", end_color="00CC00", fill_type="solid")
    fill_2nd = PatternFill(start_color="99FF99", end_color="99FF99", fill_type="solid")
    fill_3rd = PatternFill(start_color="E6FFE6", end_color="E6FFE6", fill_type="solid")

    # 遍历每一行（跳过表头）
    for row_idx in range(start_row + 1, end_row + 1):
        # 获取该行的数值和对应列索引
        values = []
        for col_idx in range(start_col, end_col + 1):
            cell_value = sheet.cell(row=row_idx, column=col_idx).value
            if isinstance(cell_value, (int, float)) and not pd.isna(cell_value):
                values.append((cell_value, col_idx))

        # 按数值降序排序
        values.sort(key=lambda x: x[0], reverse=True)

        # 为前三名标色
        for rank, (val, col_idx) in enumerate(values[:3]):
            cell = sheet.cell(row=row_idx, column=col_idx)
            if rank == 0:
                cell.fill = fill_1st
            elif rank == 1:
                cell.fill = fill_2nd
            elif rank == 2:
                cell.fill = fill_3rd
            # 加粗字体
            cell.font = Font(bold=True)


class GAT(nn.Module):
    def __init__(self, in_channels, out_channels):
        super().__init__()
        self.dropout = FIXED_DROPOUT
        self.conv1 = GATConv(in_channels, FIXED_GAT_HIDDEN_DIM, heads=FIXED_GAT_HEADS[0], dropout=self.dropout)
        self.conv2 = GATConv(FIXED_GAT_HIDDEN_DIM * FIXED_GAT_HEADS[0], out_channels, heads=FIXED_GAT_HEADS[1],
                             concat=False, dropout=self.dropout)

    def forward(self, x, edge_index):
        x = F.dropout(x, p=self.dropout, training=self.training)
        x1 = F.elu(self.conv1(x, edge_index))
        x = F.dropout(x1, p=self.dropout, training=self.training)
        x2 = self.conv2(x, edge_index)
        return x2, x1

# ---------- PreGS相关模块 ----------
class GraphSAGE(nn.Module):
    def __init__(self, in_feats, aggr):
        super(GraphSAGE, self).__init__()
        self.conv1 = SAGEConv(in_feats, FIXED_GAT_HIDDEN_DIM, aggr=aggr, bias=False, root_weight=False)
        self.dropout = FIXED_DROPOUT
        self.reset_parameters()

    def reset_parameters(self):
        for m in self.modules():
            if isinstance(m, SAGEConv):
                m.reset_parameters()

    def forward(self, x, edge_index):
        x0 = F.dropout(x, p=self.dropout, training=self.training)
        x0 = F.relu(self.conv1(x0, edge_index))
        return x0


class WeightedSum(nn.Module):
    def __init__(self, num_features):
        super().__init__()
        self.weights = nn.Parameter(torch.ones(num_features, device=device))
        self.softmax = nn.Softmax(dim=0)

    def forward(self, features):
        normalized_weights = self.softmax(self.weights)
        weighted_sum = sum(w * feat for w, feat in zip(normalized_weights, features))
        return weighted_sum


# ---------- PreGS完整模型 ----------
class PreGS(nn.Module):
    """
    PreGS完整模型固定使用以下三类输入：
    1. 多个GraphSAGE专家的加权融合特征；
    2. GAT第一层多头特征的加权融合表示；
    3. 节点原始特征。

    MLP分类结果最后与预训练GAT的最终logits进行可学习加权融合。
    """

    def __init__(
        self,
        in_dim,
        raw_dim,
        num_classes
    ):
        super().__init__()

        # GS融合表示和GAT第一层融合表示的维度均为in_dim，
        # 再拼接raw_dim维的原始节点特征。
        mlp_in_dim = in_dim + in_dim + raw_dim

        self.fc1 = nn.Linear(
            mlp_in_dim,
            FIXED_MLP_HIDDEN_DIM
        )

        self.fc2 = nn.Linear(
            FIXED_MLP_HIDDEN_DIM,
            num_classes
        )

        self.dropout = nn.Dropout(
            FIXED_DROPOUT
        )

        # PreGS分类结果与预训练GAT最终输出的融合权重
        self.fusion_weight = nn.Parameter(
            torch.ones(2, device=device)
        )

        self.softmax = nn.Softmax(dim=0)

        self.reset_parameters()

    def reset_parameters(self):
        for module in self.modules():
            if isinstance(module, nn.Linear):
                nn.init.xavier_uniform_(
                    module.weight
                )

                if module.bias is not None:
                    nn.init.zeros_(
                        module.bias
                    )

        nn.init.ones_(
            self.fusion_weight
        )

    def forward(
        self,
        gs_feat,
        gat1_feat,
        raw_feat,
        gat_final_output
    ):
        # 完整模型固定拼接三类特征
        fused_input = torch.cat(
            [
                gs_feat,
                gat1_feat,
                raw_feat
            ],
            dim=1
        )

        hidden = F.relu(
            self.fc1(fused_input)
        )

        hidden = self.dropout(hidden)

        pregs_output = self.fc2(hidden)

        # 与预训练GAT最终输出进行可学习加权融合
        normalized_weight = self.softmax(
            self.fusion_weight
        )

        output = (
            normalized_weight[0] * pregs_output
            + normalized_weight[1] * gat_final_output
        )

        return output

class SourceWeightedConcat(nn.Module):
    """
    对原始特征、GS融合特征和GAT融合特征进行源级加权拼接。
    完整PreGSv2固定启用可学习源级权重。
    """

    def __init__(self, num_sources=3):
        super().__init__()

        self.weights = nn.Parameter(
            torch.ones(
                num_sources,
                device=device
            )
        )

    def forward(
        self,
        raw_feat,
        gs_feat,
        gat_feat
    ):
        normalized_weights = torch.softmax(
            self.weights,
            dim=0
        )

        return torch.cat(
            [
                normalized_weights[0] * raw_feat,
                normalized_weights[1] * gs_feat,
                normalized_weights[2] * gat_feat
            ],
            dim=1
        )

class GatedPreGSv2(nn.Module):
    """
    PreGSv2完整模型：

    1. 对raw、GS融合表示和GAT融合表示进行源级加权拼接；
    2. 使用GAT第一层融合表示生成结构门控；
    3. 使用门控后的融合表示进行MLP分类；
    4. MLP分类结果与预训练GAT最终logits进行加权融合。
    """

    def __init__(
        self,
        raw_dim,
        hidden_dim,
        num_classes
    ):
        super().__init__()

        self.in_dim = (
            raw_dim
            + hidden_dim
            + hidden_dim
        )

        # 完整PreGSv2固定启用结构门控
        self.gate_layer = nn.Linear(
            hidden_dim,
            self.in_dim
        )

        self.fc1 = nn.Linear(
            self.in_dim,
            FIXED_MLP_HIDDEN_DIM
        )

        self.fc2 = nn.Linear(
            FIXED_MLP_HIDDEN_DIM,
            num_classes
        )

        self.dropout = nn.Dropout(
            FIXED_DROPOUT
        )

        self.fusion_weight = nn.Parameter(
            torch.ones(2, device=device)
        )

        self.softmax = nn.Softmax(dim=0)

        self.reset_parameters()

    def reset_parameters(self):
        for module in self.modules():
            if isinstance(module, nn.Linear):
                nn.init.xavier_uniform_(
                    module.weight
                )

                if module.bias is not None:
                    nn.init.zeros_(
                        module.bias
                    )

        nn.init.ones_(
            self.fusion_weight
        )

    def forward(
        self,
        fused_feat,
        gate_feature,
        gat_final_output
    ):
        # 使用GAT融合表示生成逐维门控向量
        gate_vector = torch.sigmoid(
            self.gate_layer(
                gate_feature
            )
        )

        gated_feature = (
            fused_feat * gate_vector
        )

        hidden = F.relu(
            self.fc1(gated_feature)
        )

        hidden = self.dropout(hidden)

        pregs_v2_output = self.fc2(hidden)

        normalized_weight = self.softmax(
            self.fusion_weight
        )

        output = (
            normalized_weight[0]
            * pregs_v2_output
            + normalized_weight[1]
            * gat_final_output
        )

        return output

# ---------- 工具函数 ----------
def build_feature_list(gs_features):
    final_features = []
    if gs_features is not None and len(gs_features) == 8:
        final_features.extend(gs_features)
    return final_features


# ---------- 通用训练函数 ----------
def train_baseline_model(
    model,
    data,
    run_seed
):
    set_random_seed(run_seed)

    model.to(device).train()
    optimizer = torch.optim.Adam(model.parameters(), lr=FIXED_LR, weight_decay=FIXED_WD)
    criterion = torch.nn.CrossEntropyLoss()

    best_val_acc = 0.0
    curr_patience = 0
    best_state_dict = None

    x = data.x.to(device)
    edge_index = data.edge_index.to(device)
    train_mask = data.train_mask.to(device)
    val_mask = data.val_mask.to(device)
    test_mask = data.test_mask.to(device)
    y = data.y.to(device)

    for epoch in range(FIXED_EPOCHS):
        optimizer.zero_grad()

        logits, _ = model(x, edge_index)

        loss = criterion(logits[train_mask], y[train_mask])
        loss.backward()
        optimizer.step()

        model.eval()
        with torch.no_grad():
            val_logits, _ = model(x, edge_index)

            val_pred = val_logits[val_mask].argmax(dim=1)
            val_acc = (val_pred == y[val_mask]).float().mean().item()

        if val_acc > best_val_acc:
            best_val_acc = val_acc
            best_state_dict = {k: v.cpu().clone() for k, v in model.state_dict().items()}
            curr_patience = 0
        else:
            curr_patience += 1

        if curr_patience >= FIXED_PATIENCE:
            break

        model.train()

    if best_state_dict is not None:
        model.load_state_dict(best_state_dict)
    model.eval()

    with torch.no_grad():
        test_logits, _ = model(x, edge_index)

        test_pred = test_logits[test_mask].argmax(dim=1)
        test_acc = (test_pred == y[test_mask]).float().mean().item()
        test_f1 = compute_macro_f1(y[test_mask], test_pred)

    return best_val_acc, test_acc, test_f1, model

# ---------- 训练PreGS完整模型 ----------
def train_pregs(
    data,
    trained_gat,
    run_seed
):
    set_random_seed(run_seed)

    num_features = data.x.shape[1]
    num_classes = len(
        torch.unique(data.y)
    )

    head_dim = FIXED_GAT_HIDDEN_DIM

    x = data.x.to(device)

    edge_index_with_self = (
        add_remaining_self_loops(
            data.edge_index
        )[0].to(device)
    )

    # ========================================================
    # 1. 根据预训练GAT第一层参数构建GraphSAGE专家
    # ========================================================
    feature_extractors = []

    for aggregator in FIXED_GS_AGGRS:
        graph_sage_expert = GraphSAGE(
            num_features,
            aggregator
        ).to(device)

        feature_extractors.append(
            graph_sage_expert
        )

    gat_first_layer_weight = (
        trained_gat.conv1.lin.weight.detach()
    )

    gat_head_weights = [
        gat_first_layer_weight[
            head_index * head_dim:
            (head_index + 1) * head_dim,
            :
        ]
        for head_index in range(
            FIXED_GAT_HEADS[0]
        )
    ]

    for expert_index, graph_sage_expert in enumerate(
        feature_extractors
    ):
        with torch.no_grad():
            graph_sage_expert.conv1.lin_l.weight.copy_(
                gat_head_weights[
                    expert_index
                    % len(gat_head_weights)
                ]
            )

            for parameter in graph_sage_expert.parameters():
                parameter.requires_grad = False

        graph_sage_expert.eval()

    # ========================================================
    # 2. 提取固定的GS专家特征和GAT特征
    # ========================================================
    trained_gat.eval()

    with torch.no_grad():
        gs_features = [
            graph_sage_expert(
                x,
                edge_index_with_self
            )
            for graph_sage_expert
            in feature_extractors
        ]

        final_gs_features = build_feature_list(
            gs_features
        )

        (
            gat_final_output,
            gat_first_layer_output
        ) = trained_gat(
            x,
            edge_index_with_self
        )

        gat_first_layer_features = torch.split(
            gat_first_layer_output,
            head_dim,
            dim=1
        )

        raw_feature = x

    # ========================================================
    # 3. 可学习的专家融合模块
    # ========================================================
    gs_sum_module = WeightedSum(
        len(final_gs_features)
    ).to(device)

    gat_sum_module = WeightedSum(
        len(gat_first_layer_features)
    ).to(device)

    pregs_model = PreGS(
        in_dim=FIXED_GAT_HIDDEN_DIM,
        raw_dim=num_features,
        num_classes=num_classes
    ).to(device)

    optimizer = torch.optim.Adam(
        (
            list(pregs_model.parameters())
            + list(gs_sum_module.parameters())
            + list(gat_sum_module.parameters())
        ),
        lr=FIXED_LR,
        weight_decay=FIXED_WD
    )

    criterion = nn.CrossEntropyLoss()

    train_mask = data.train_mask.to(device)
    val_mask = data.val_mask.to(device)
    test_mask = data.test_mask.to(device)
    labels = data.y.to(device)

    best_val_acc = 0.0
    best_test_acc = 0.0
    best_test_f1 = 0.0
    current_patience = 0
    best_state = None

    start_time = time.time()

    # ========================================================
    # 4. 训练融合模块和分类器
    # ========================================================
    for epoch in range(FIXED_EPOCHS):
        pregs_model.train()
        gs_sum_module.train()
        gat_sum_module.train()

        optimizer.zero_grad()

        gs_fused_feature = gs_sum_module(
            final_gs_features
        )

        gat_fused_feature = gat_sum_module(
            gat_first_layer_features
        )

        logits = pregs_model(
            gs_feat=gs_fused_feature,
            gat1_feat=gat_fused_feature,
            raw_feat=raw_feature,
            gat_final_output=gat_final_output
        )

        loss = criterion(
            logits[train_mask],
            labels[train_mask]
        )

        loss.backward()
        optimizer.step()

        # ----------------------------------------------------
        # 验证与当前测试结果
        # ----------------------------------------------------
        pregs_model.eval()
        gs_sum_module.eval()
        gat_sum_module.eval()

        with torch.no_grad():
            gs_fused_feature_eval = (
                gs_sum_module(
                    final_gs_features
                )
            )

            gat_fused_feature_eval = (
                gat_sum_module(
                    gat_first_layer_features
                )
            )

            evaluation_logits = pregs_model(
                gs_feat=gs_fused_feature_eval,
                gat1_feat=gat_fused_feature_eval,
                raw_feat=raw_feature,
                gat_final_output=gat_final_output
            )

            val_prediction = (
                evaluation_logits[val_mask]
                .argmax(dim=1)
            )

            val_acc = (
                val_prediction
                == labels[val_mask]
            ).float().mean().item()

            test_prediction = (
                evaluation_logits[test_mask]
                .argmax(dim=1)
            )

            test_acc = (
                test_prediction
                == labels[test_mask]
            ).float().mean().item()

            test_f1 = compute_macro_f1(
                labels[test_mask],
                test_prediction
            )

        if val_acc > best_val_acc:
            best_val_acc = val_acc
            best_test_acc = test_acc
            best_test_f1 = test_f1
            current_patience = 0

            best_state = {
                'model': {
                    key: value.detach().cpu().clone()
                    for key, value
                    in pregs_model.state_dict().items()
                },
                'gs_sum': {
                    key: value.detach().cpu().clone()
                    for key, value
                    in gs_sum_module.state_dict().items()
                },
                'gat_sum': {
                    key: value.detach().cpu().clone()
                    for key, value
                    in gat_sum_module.state_dict().items()
                }
            }

        else:
            current_patience += 1

            if current_patience >= FIXED_PATIENCE:
                break

    # ========================================================
    # 5. 恢复验证集表现最优的参数
    # ========================================================
    if best_state is not None:
        pregs_model.load_state_dict(
            best_state['model']
        )

        gs_sum_module.load_state_dict(
            best_state['gs_sum']
        )

        gat_sum_module.load_state_dict(
            best_state['gat_sum']
        )

    train_duration = (
        time.time() - start_time
    )

    return (
        best_test_acc,
        best_test_f1,
        train_duration,
        best_val_acc
    )

# ---------- 训练PreGSv2完整模型 ----------
def train_pregs_v2(
    data,
    trained_gat,
    run_seed
):
    set_random_seed(run_seed)

    num_features = data.x.shape[1]
    num_classes = len(torch.unique(data.y))
    head_dim = FIXED_GAT_HIDDEN_DIM

    x = data.x.to(device)
    edge_index_with_self = add_remaining_self_loops(data.edge_index)[0].to(device)

    feature_extractors = []
    for aggr in FIXED_GS_AGGRS:
        gs = GraphSAGE(num_features, aggr).to(device)
        feature_extractors.append(gs)

    gat1_w = trained_gat.conv1.lin.weight.detach()
    gat1_heads = [gat1_w[i * head_dim: (i + 1) * head_dim, :] for i in range(FIXED_GAT_HEADS[0])]

    for h, gs in enumerate(feature_extractors):
        with torch.no_grad():
            gs.conv1.lin_l.weight.copy_(gat1_heads[h % len(gat1_heads)])
            for param in gs.parameters():
                param.requires_grad = False
        gs.eval()

    trained_gat.eval()
    with torch.no_grad():
        gs_features = [gs(x, edge_index_with_self) for gs in feature_extractors]
        gat_final_output, gat1_output = trained_gat(x, edge_index_with_self)
        gat1_features = torch.split(gat1_output, head_dim, dim=1)
        raw_feat = x

    gs_sum_module = WeightedSum(len(gs_features)).to(device)
    gat1_sum_module = WeightedSum(len(gat1_features)).to(device)
    source_concat_module = SourceWeightedConcat(
        num_sources=3
    ).to(device)

    model = GatedPreGSv2(
        raw_dim=num_features,
        hidden_dim=FIXED_GAT_HIDDEN_DIM,
        num_classes=num_classes
    ).to(device)

    optimizer = torch.optim.Adam(
        list(model.parameters())
        + list(gs_sum_module.parameters())
        + list(gat1_sum_module.parameters())
        + list(source_concat_module.parameters()),
        lr=FIXED_LR,
        weight_decay=FIXED_WD
    )
    criterion = nn.CrossEntropyLoss()

    best_val_acc = 0.0
    best_test_acc = 0.0
    best_test_f1 = 0.0
    curr_patience = 0
    best_state = None
    start_time = time.time()

    train_mask = data.train_mask.to(device)
    val_mask = data.val_mask.to(device)
    test_mask = data.test_mask.to(device)
    y = data.y.to(device)

    for epoch in range(FIXED_EPOCHS):
        model.train()
        gs_sum_module.train()
        gat1_sum_module.train()
        source_concat_module.train()
        optimizer.zero_grad()

        gs_fused = gs_sum_module(gs_features)
        gat_fused = gat1_sum_module(gat1_features)
        fused_feat = source_concat_module(raw_feat, gs_fused, gat_fused)
        logits = model(fused_feat, gat_fused, gat_final_output)

        loss = criterion(logits[train_mask], y[train_mask])
        loss.backward()
        optimizer.step()

        model.eval()
        gs_sum_module.eval()
        gat1_sum_module.eval()
        source_concat_module.eval()
        with torch.no_grad():
            gs_fused_val = gs_sum_module(gs_features)
            gat_fused_val = gat1_sum_module(gat1_features)
            fused_feat_val = source_concat_module(raw_feat, gs_fused_val, gat_fused_val)
            val_logits = model(fused_feat_val, gat_fused_val, gat_final_output)

            val_pred = val_logits[val_mask].argmax(dim=1)
            val_acc = (val_pred == y[val_mask]).float().mean().item()
            test_pred = val_logits[test_mask].argmax(dim=1)
            test_acc = (test_pred == y[test_mask]).float().mean().item()
            test_f1 = compute_macro_f1(y[test_mask], test_pred)

        if val_acc > best_val_acc:
            best_val_acc = val_acc
            best_test_acc = test_acc
            best_test_f1 = test_f1
            curr_patience = 0

            best_state = {
                'model': {k: v.detach().cpu().clone() for k, v in model.state_dict().items()},
                'gs_sum': {k: v.detach().cpu().clone() for k, v in gs_sum_module.state_dict().items()},
                'gat_sum': {k: v.detach().cpu().clone() for k, v in gat1_sum_module.state_dict().items()},
                'source_concat': {k: v.detach().cpu().clone() for k, v in source_concat_module.state_dict().items()}
            }
        else:
            curr_patience += 1
            if curr_patience >= FIXED_PATIENCE:
                break

    if best_state is not None:
        model.load_state_dict(best_state['model'])
        gs_sum_module.load_state_dict(best_state['gs_sum'])
        gat1_sum_module.load_state_dict(best_state['gat_sum'])
        source_concat_module.load_state_dict(best_state['source_concat'])

    train_duration = time.time() - start_time
    return best_test_acc, best_test_f1, train_duration, best_val_acc


# ---------- 实验运行函数 ----------
def run_dataset_percent_experiment(
    dataset_name,
    data,
    train_percent
):
    print("\n\n" + "#" * 100)
    print(
        f"######################## "
        f"{dataset_name} "
        f"(训练集比例 {train_percent}) "
        f"########################"
    )

    num_nodes = data.num_nodes
    num_features = data.x.shape[1]
    num_classes = len(torch.unique(data.y))

    print(
        f"数据集信息 - "
        f"节点数: {num_nodes}, "
        f"特征维度: {num_features}, "
        f"类别数: {num_classes}"
    )

    train_size = int(train_percent * num_nodes)
    val_size = int(VAL_PERCENT * num_nodes)

    # 测试阶段为2，正式实验时恢复为30
    repeat_times = 2

    # 只取得本次实际需要运行的模型
    active_models = get_active_models()

    # 只为被选中的模型创建结果容器
    models_results = {
        model_name: []
        for model_name in active_models
    }

    print(
        "本次运行模型："
        + ", ".join(active_models)
    )

    for repeat_index in range(repeat_times):
        print("\n" + "-" * 80)
        print(
            f"{dataset_name} "
            f"(训练集比例 {train_percent}) - "
            f"重复实验 "
            f"{repeat_index + 1}/{repeat_times}"
        )
        print("-" * 80 + "\n")

        repeat_seed = seed + repeat_index * 100

        # ----------------------------------------------------
        # 构造当前重复实验的数据划分
        # ----------------------------------------------------
        set_random_seed(repeat_seed)

        indices = torch.randperm(num_nodes)

        train_mask = torch.zeros(
            num_nodes,
            dtype=torch.bool
        )
        val_mask = torch.zeros(
            num_nodes,
            dtype=torch.bool
        )
        test_mask = torch.zeros(
            num_nodes,
            dtype=torch.bool
        )

        train_mask[
            indices[:train_size]
        ] = True

        val_mask[
            indices[
                train_size:
                train_size + val_size
            ]
        ] = True

        test_mask[
            indices[
                train_size + val_size:
            ]
        ] = True

        data.train_mask = train_mask.to(device)
        data.val_mask = val_mask.to(device)
        data.test_mask = test_mask.to(device)

        # 每次重复实验都重新训练一个 GAT。
        # 当选择 PreGS 或 PreGSv2 时，后续直接复用它。
        trained_gat = None

        total_models = len(active_models)

        # ----------------------------------------------------
        # 只循环运行 SELECTED_MODELS 中的模型
        # ----------------------------------------------------
        for model_index, model_name in enumerate(
            active_models,
            start=1
        ):
            print(
                f"\n[{model_index}/{total_models}] "
                f"训练 {model_name}"
            )

            # 在每个模型创建前重置随机种子。
            # 这样改变 SELECTED_MODELS 不会改变其他模型的初始化。
            set_random_seed(repeat_seed)

            result, trained_gat = run_one_selected_model(
                model_name=model_name,
                data=data,
                num_features=num_features,
                num_classes=num_classes,
                trained_gat=trained_gat,
                run_seed=repeat_seed
            )

            models_results[model_name].append(result)

            print(
                f"{model_name} - "
                f"测试准确率: {result['acc']:.5f}, "
                f"Macro-F1: {result['f1']:.5f}, "
                f"验证准确率: {result['val_acc']:.5f}, "
                f"耗时: {result['time']:.2f}s"
            )

    # --------------------------------------------------------
    # 汇总本数据集、本训练比例下的重复实验结果
    # --------------------------------------------------------
    for model_name in active_models:
        result_list = models_results[model_name]

        if not result_list:
            continue

        accs = [
            result['acc']
            for result in result_list
        ]

        f1s = [
            result['f1']
            for result in result_list
        ]

        times = [
            result['time']
            for result in result_list
        ]

        val_accs = [
            result['val_acc']
            for result in result_list
        ]

        all_results_raw.append({
            'Dataset': dataset_name,
            'TrainPercent': train_percent,
            'Model': model_name,

            'MeanTestAcc': float(
                np.mean(accs)
            ),
            'StdTestAcc': compute_sample_std(
                accs
            ),

            'MeanTestF1': float(
                np.mean(f1s)
            ),
            'StdTestF1': compute_sample_std(
                f1s
            ),

            'MeanTime': float(
                np.mean(times)
            ),
            'StdTime': compute_sample_std(
                times
            ),

            'MeanValAcc': float(
                np.mean(val_accs)
            ),
            'StdValAcc': compute_sample_std(
                val_accs
            ),

            'NumRepeats': len(result_list),

            'FixedParams': (
                f"lr={FIXED_LR}, "
                f"wd={FIXED_WD}, "
                f"dropout={FIXED_DROPOUT}"
            )
        })

    if torch.cuda.is_available():
        torch.cuda.empty_cache()

def run_one_selected_model(
    model_name,
    data,
    num_features,
    num_classes,
    trained_gat,
    run_seed
):
    """
    运行一个被选中的模型。

    返回：
    1. 当前模型的实验结果字典；
    2. 已训练的 GAT 模型。
    """

    # 确保模型构造前使用当前重复实验的随机种子
    set_random_seed(run_seed)


    # --------------------------------------------------------
    # GAT 预训练
    # --------------------------------------------------------
    if model_name == MODEL_GAT:
        model = GAT(
            num_features,
            num_classes
        )

        start_time = time.time()

        val_acc, test_acc, test_f1, trained_gat = (
            train_baseline_model(
                model,
                data,
                run_seed
            )
        )

        duration = time.time() - start_time

        result = {
            'acc': test_acc,
            'f1': test_f1,
            'time': duration,
            'val_acc': val_acc,
        }

        return result, trained_gat

    # --------------------------------------------------------
    # PreGS 完整模型
    # --------------------------------------------------------
    if model_name == MODEL_PREGS:
        if trained_gat is None:
            raise RuntimeError(
                "运行PreGS前没有可用的trained_gat。"
                "请确保SELECTED_MODELS中包含MODEL_GAT。"
            )

        (
            test_acc,
            test_f1,
            duration,
            val_acc
        ) = train_pregs(
            data,
            trained_gat,
            run_seed
        )

        result = {
            'acc': test_acc,
            'f1': test_f1,
            'time': duration,
            'val_acc': val_acc,
        }

        return result, trained_gat

    # --------------------------------------------------------
    # PreGSv2 完整模型
    # --------------------------------------------------------
    if model_name == MODEL_PREGS_V2:
        if trained_gat is None:
            raise RuntimeError(
                "运行 PreGSv2 前没有可用的 trained_gat。"
                "请确保 SELECTED_MODELS 中包含 MODEL_GAT。"
            )

        (
            test_acc,
            test_f1,
            duration,
            val_acc
        ) = train_pregs_v2(
            data,
            trained_gat,
            run_seed
        )

        result = {
            'acc': test_acc,
            'f1': test_f1,
            'time': duration,
            'val_acc': val_acc,
        }

        return result, trained_gat

    raise ValueError(
        f"尚未定义模型 {model_name} 的运行方式。"
    )

# ---------- 数据集加载 ----------
def load_dataset(dataset_name):
    base_path = os.path.join(
        os.path.dirname(os.path.abspath(__file__)),
        "datasets",
        dataset_name
    )
    if not os.path.exists(base_path):
        print(f"数据集 {dataset_name} 不存在")
        return None
    try:
        adj = np.load(os.path.join(base_path, f"{dataset_name}_adj.npy"))
        feat = np.load(os.path.join(base_path, f"{dataset_name}_feat.npy"))
        label = np.load(os.path.join(base_path, f"{dataset_name}_label.npy"))
        edge_index = np.array(np.nonzero(adj))
        data = Data(
            x=torch.tensor(feat, dtype=torch.float),
            edge_index=torch.tensor(edge_index, dtype=torch.long),
            y=torch.tensor(label, dtype=torch.long)
        )
        return data
    except Exception as e:
        print(f"加载 {dataset_name} 失败: {e}")
        return None


# ---------- 增量保存实验结果 ----------
def save_results_to_excel(results, out_file):
    """
    将当前累计实验结果保存到同一个 Excel 文件。

    调用时机：
    每完成一个数据集的全部训练比例和重复实验后调用一次。

    保存策略：
    1. 先写入临时文件；
    2. 临时文件完整生成后，再替换正式结果文件；
    3. 避免写入过程中断导致原结果文件损坏。
    """
    if not results:
        return None

    # ============================================================
    # 1. 整理原始结果
    # ============================================================
    df_raw = pd.DataFrame(results).copy()

    df_raw['Model'] = pd.Categorical(
        df_raw['Model'],
        categories=MODEL_ORDER,
        ordered=True
    )

    df_raw = df_raw.sort_values(
        ['Dataset', 'TrainPercent', 'Model']
    )

    # 本次累计结果中实际出现的模型
    present_models = [
        model_name
        for model_name in MODEL_ORDER
        if (df_raw['Model'] == model_name).any()
    ]

    # ============================================================
    # 2. 生成百分比及“均值 ± 标准差”结果
    # ============================================================
    df_read = df_raw.copy()

    df_read['MeanTestAcc(%)'] = (
        df_read['MeanTestAcc'] * 100
    )
    df_read['StdTestAcc(%)'] = (
        df_read['StdTestAcc'] * 100
    )

    df_read['MeanTestF1(%)'] = (
        df_read['MeanTestF1'] * 100
    )
    df_read['StdTestF1(%)'] = (
        df_read['StdTestF1'] * 100
    )

    df_read['MeanValAcc(%)'] = (
        df_read['MeanValAcc'] * 100
    )
    df_read['StdValAcc(%)'] = (
        df_read['StdValAcc'] * 100
    )

    df_read['TestAcc_MeanStd(%)'] = [
        format_mean_std(
            mean_value=mean_value,
            std_value=std_value,
            scale=100,
            decimals=2
        )
        for mean_value, std_value in zip(
            df_read['MeanTestAcc'],
            df_read['StdTestAcc']
        )
    ]

    df_read['TestF1_MeanStd(%)'] = [
        format_mean_std(
            mean_value=mean_value,
            std_value=std_value,
            scale=100,
            decimals=2
        )
        for mean_value, std_value in zip(
            df_read['MeanTestF1'],
            df_read['StdTestF1']
        )
    ]

    df_read['ValAcc_MeanStd(%)'] = [
        format_mean_std(
            mean_value=mean_value,
            std_value=std_value,
            scale=100,
            decimals=2
        )
        for mean_value, std_value in zip(
            df_read['MeanValAcc'],
            df_read['StdValAcc']
        )
    ]

    df_read['Time_MeanStd(s)'] = [
        format_mean_std(
            mean_value=mean_value,
            std_value=std_value,
            scale=1,
            decimals=2
        )
        for mean_value, std_value in zip(
            df_read['MeanTime'],
            df_read['StdTime']
        )
    ]

    # ============================================================
    # 3. 构建均值对比表
    # ============================================================
    pivot_acc = df_read.pivot_table(
        index=['Dataset', 'TrainPercent'],
        columns='Model',
        values='MeanTestAcc(%)',
        aggfunc='mean',
        observed=True
    ).round(2)

    pivot_f1 = df_read.pivot_table(
        index=['Dataset', 'TrainPercent'],
        columns='Model',
        values='MeanTestF1(%)',
        aggfunc='mean',
        observed=True
    ).round(2)

    pivot_acc = pivot_acc.reindex(
        columns=present_models
    )

    pivot_f1 = pivot_f1.reindex(
        columns=present_models
    )

    # ============================================================
    # 4. 构建“均值 ± 标准差”对比表
    # ============================================================
    pivot_acc_mean_std = df_read.pivot_table(
        index=['Dataset', 'TrainPercent'],
        columns='Model',
        values='TestAcc_MeanStd(%)',
        aggfunc='first',
        observed=True
    ).reindex(
        columns=present_models
    )

    pivot_f1_mean_std = df_read.pivot_table(
        index=['Dataset', 'TrainPercent'],
        columns='Model',
        values='TestF1_MeanStd(%)',
        aggfunc='first',
        observed=True
    ).reindex(
        columns=present_models
    )

    pivot_val_mean_std = df_read.pivot_table(
        index=['Dataset', 'TrainPercent'],
        columns='Model',
        values='ValAcc_MeanStd(%)',
        aggfunc='first',
        observed=True
    ).reindex(
        columns=present_models
    )

    pivot_time_mean_std = df_read.pivot_table(
        index=['Dataset', 'TrainPercent'],
        columns='Model',
        values='Time_MeanStd(s)',
        aggfunc='first',
        observed=True
    ).reindex(
        columns=present_models
    )

    # ============================================================
    # 5. 计算最优模型与PreGS相对GAT提升
    # ============================================================
    def get_available_models_from_row(row):
        return [
            model_name
            for model_name in present_models
            if (
                model_name in row.index
                and pd.notna(row[model_name])
            )
        ]

    def best_row_acc(row):
        available_models = get_available_models_from_row(
            row
        )

        if not available_models:
            return pd.Series(
                [np.nan, np.nan, None, np.nan],
                index=[
                    'PreGS-GAT提升(%)',
                    'PreGSv2-GAT提升(%)',
                    '最优模型',
                    '最优准确率(%)'
                ]
            )

        best_model = max(
            available_models,
            key=lambda model_name: row[model_name]
        )

        if (
            MODEL_PREGS in available_models
            and MODEL_GAT in available_models
        ):
            pregs_gain = round(
                row[MODEL_PREGS] - row[MODEL_GAT],
                2
            )
        else:
            pregs_gain = np.nan

        if (
            MODEL_PREGS_V2 in available_models
            and MODEL_GAT in available_models
        ):
            pregs_v2_gain = round(
                row[MODEL_PREGS_V2] - row[MODEL_GAT],
                2
            )
        else:
            pregs_v2_gain = np.nan

        return pd.Series(
            [
                pregs_gain,
                pregs_v2_gain,
                best_model,
                round(row[best_model], 2)
            ],
            index=[
                'PreGS-GAT提升(%)',
                'PreGSv2-GAT提升(%)',
                '最优模型',
                '最优准确率(%)'
            ]
        )

    def best_row_f1(row):
        available_models = get_available_models_from_row(
            row
        )

        if not available_models:
            return pd.Series(
                [np.nan, np.nan, None, np.nan],
                index=[
                    'PreGS-GAT提升(%)',
                    'PreGSv2-GAT提升(%)',
                    '最优模型',
                    '最优F1(%)'
                ]
            )

        best_model = max(
            available_models,
            key=lambda model_name: row[model_name]
        )

        if (
            MODEL_PREGS in available_models
            and MODEL_GAT in available_models
        ):
            pregs_gain = round(
                row[MODEL_PREGS] - row[MODEL_GAT],
                2
            )
        else:
            pregs_gain = np.nan

        if (
            MODEL_PREGS_V2 in available_models
            and MODEL_GAT in available_models
        ):
            pregs_v2_gain = round(
                row[MODEL_PREGS_V2] - row[MODEL_GAT],
                2
            )
        else:
            pregs_v2_gain = np.nan

        return pd.Series(
            [
                pregs_gain,
                pregs_v2_gain,
                best_model,
                round(row[best_model], 2)
            ],
            index=[
                'PreGS-GAT提升(%)',
                'PreGSv2-GAT提升(%)',
                '最优模型',
                '最优F1(%)'
            ]
        )

    pivot_acc[
        [
            'PreGS-GAT提升(%)',
            'PreGSv2-GAT提升(%)',
            '最优模型',
            '最优准确率(%)'
        ]
    ] = pivot_acc.apply(
        best_row_acc,
        axis=1
    )

    pivot_f1[
        [
            'PreGS-GAT提升(%)',
            'PreGSv2-GAT提升(%)',
            '最优模型',
            '最优F1(%)'
        ]
    ] = pivot_f1.apply(
        best_row_f1,
        axis=1
    )

    # ============================================================
    # 6. 先写入临时Excel，再安全替换正式文件
    # ============================================================
    file_root, file_extension = os.path.splitext(
        out_file
    )

    temp_file = (
        f"{file_root}_temp{file_extension}"
    )

    if os.path.exists(temp_file):
        os.remove(temp_file)

    try:
        with pd.ExcelWriter(
            temp_file,
            engine='openpyxl'
        ) as writer:

            df_raw.to_excel(
                writer,
                sheet_name='基础模型_原始',
                index=False
            )

            df_read.to_excel(
                writer,
                sheet_name='基础模型_百分比',
                index=False
            )

            pivot_acc.to_excel(
                writer,
                sheet_name='准确率均值对比'
            )

            pivot_f1.to_excel(
                writer,
                sheet_name='MacroF1均值对比'
            )

            pivot_acc_mean_std.to_excel(
                writer,
                sheet_name='准确率_均值标准差'
            )

            pivot_f1_mean_std.to_excel(
                writer,
                sheet_name='MacroF1_均值标准差'
            )

            pivot_val_mean_std.to_excel(
                writer,
                sheet_name='验证准确率_均值标准差'
            )

            pivot_time_mean_std.to_excel(
                writer,
                sheet_name='时间_均值标准差'
            )

            # 准确率前三名标色
            ws_acc = writer.sheets[
                '准确率均值对比'
            ]

            acc_model_columns = [
                model_name
                for model_name in present_models
                if model_name in pivot_acc.columns
            ]

            if acc_model_columns:
                start_col = 3
                end_col = (
                    start_col
                    + len(acc_model_columns)
                    - 1
                )

                highlight_top3(
                    ws_acc,
                    start_row=1,
                    start_col=start_col,
                    end_row=ws_acc.max_row,
                    end_col=end_col
                )

            # Macro-F1前三名标色
            ws_f1 = writer.sheets[
                'MacroF1均值对比'
            ]

            f1_model_columns = [
                model_name
                for model_name in present_models
                if model_name in pivot_f1.columns
            ]

            if f1_model_columns:
                start_col = 3
                end_col = (
                    start_col
                    + len(f1_model_columns)
                    - 1
                )

                highlight_top3(
                    ws_f1,
                    start_row=1,
                    start_col=start_col,
                    end_row=ws_f1.max_row,
                    end_col=end_col
                )

        # 临时文件完整生成后，再覆盖正式结果文件
        os.replace(
            temp_file,
            out_file
        )

    except Exception:
        if os.path.exists(temp_file):
            os.remove(temp_file)
        raise

    return {
        'pivot_acc': pivot_acc,
        'pivot_f1': pivot_f1,
        'pivot_acc_mean_std': pivot_acc_mean_std,
        'pivot_f1_mean_std': pivot_f1_mean_std,
        'pivot_val_mean_std': pivot_val_mean_std,
        'pivot_time_mean_std': pivot_time_mean_std
    }

# ---------- 主函数 ----------
def main():
    active_models = get_active_models()

    print("=" * 100)
    print(" PreGS分类实验（可选模型、均值与标准差、增量保存）")
    print(" 本次运行模型：", ", ".join(active_models))
    print("=" * 100)

    # 正式数据集列表
    # dataset_list = [
    #     "acm",
    #     "amac",
    #     "amap",
    #     "dblp",
    #     "eat",
    #     "film",
    #     "pubmed",
    #     "texas",
    # ]

    # 当前测试只运行一个数据集
    dataset_list = ["acm"]

    # 输出文件只在程序开始时生成一次。
    # 后续每完成一个数据集，就更新这个同一个文件。
    timestamp = datetime.now().strftime(
        "%m%d-%H%M%S"
    )

    out_file = os.path.join(
        output_dir,
        (
            "PreGS_分类实验_可选模型_"
            "增量保存_Acc_F1_结果_"
            f"{timestamp}.xlsx"
        )
    )

    last_saved_tables = None

    for dataset_name in dataset_list:
        print("\n" + "=" * 100)
        print(f"开始处理数据集：{dataset_name}")
        print("=" * 100)

        data = load_dataset(dataset_name)

        if data is None:
            print(
                f"跳过数据集 {dataset_name}：加载失败"
            )
            continue

        # 记录运行该数据集前的结果数量
        result_count_before = len(
            all_results_raw
        )

        # 先完成当前数据集的所有训练比例
        for train_percent in TRAIN_PERCENT_LIST:
            run_dataset_percent_experiment(
                dataset_name,
                data.clone(),
                train_percent
            )

        # 只有当前数据集确实产生结果后才保存
        result_count_after = len(
            all_results_raw
        )

        if result_count_after > result_count_before:
            last_saved_tables = save_results_to_excel(
                results=all_results_raw,
                out_file=out_file
            )

            print("\n" + "-" * 100)
            print(
                f"数据集 {dataset_name} 的所有训练比例"
                f"和重复实验已完成。"
            )
            print(
                "当前累计结果已保存到：",
                out_file
            )
            print("-" * 100)

    if not all_results_raw:
        print("无有效实验结果，程序结束。")
        return

    if last_saved_tables is None:
        print("实验产生了结果，但Excel保存失败。")
        return

    # ============================================================
    # 所有数据集运行结束后的最终汇总
    # ============================================================
    pivot_acc = last_saved_tables[
        'pivot_acc'
    ]

    pivot_f1 = last_saved_tables[
        'pivot_f1'
    ]

    pivot_acc_mean_std = last_saved_tables[
        'pivot_acc_mean_std'
    ]

    pivot_f1_mean_std = last_saved_tables[
        'pivot_f1_mean_std'
    ]

    pivot_val_mean_std = last_saved_tables[
        'pivot_val_mean_std'
    ]

    pivot_time_mean_std = last_saved_tables[
        'pivot_time_mean_std'
    ]

    print("\n\n" + "=" * 100)
    print("全部实验完成！")
    print("最终结果文件：", out_file)

    print("\n==== 准确率均值汇总 ====")
    print(pivot_acc.to_string())

    print("\n==== 准确率：均值 ± 标准差 ====")
    print(pivot_acc_mean_std.to_string())

    print("\n==== Macro-F1均值汇总 ====")
    print(pivot_f1.to_string())

    print("\n==== Macro-F1：均值 ± 标准差 ====")
    print(pivot_f1_mean_std.to_string())

    print("\n==== 验证准确率：均值 ± 标准差 ====")
    print(pivot_val_mean_std.to_string())

    print("\n==== 运行时间：均值 ± 标准差 ====")
    print(pivot_time_mean_std.to_string())


if __name__ == "__main__":
    main()